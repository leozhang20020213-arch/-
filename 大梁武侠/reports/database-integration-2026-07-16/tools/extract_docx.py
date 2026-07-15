from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def clean(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split())


def extract_docx(source: Path, out_dir: Path) -> dict:
    document = Document(source)
    paragraphs = []
    headings = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = clean(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        paragraphs.append({"index": index, "style": style, "text": text})
        if style.lower().startswith("heading") or style.startswith("标题"):
            headings.append({"index": index, "style": style, "text": text})

    tables = []
    for table_index, table in enumerate(document.tables, 1):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        tables.append({
            "index": table_index,
            "rows": len(rows),
            "columns": max((len(row) for row in rows), default=0),
            "values": rows,
        })

    stem = source.stem
    lines = [f"# {stem}", "", f"来源：`{source}`", "", "## 正文段落", ""]
    for paragraph in paragraphs:
        style = paragraph["style"]
        text = paragraph["text"]
        if style.lower().startswith("heading") or style.startswith("标题"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(6, max(2, int(digits) + 1 if digits else 2))
            lines.extend(["#" * level + " " + text, ""])
        else:
            lines.extend([text, ""])

    lines.extend(["## 表格", ""])
    for table in tables:
        lines.extend([f"### 表格 {table['index']}（{table['rows']}×{table['columns']}）", ""])
        for row in table["values"]:
            lines.append(" | ".join(value.replace("|", "\\|") for value in row))
        lines.append("")

    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"{stem}.md").write_text("\n".join(lines), encoding="utf-8")
    result = {
        "source": str(source),
        "paragraphCount": len(paragraphs),
        "headingCount": len(headings),
        "headings": headings,
        "tableCount": len(tables),
        "tables": [{k: v for k, v in table.items() if k != "values"} for table in tables],
    }
    (out_dir / f"{stem}.summary.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> None:
    workspace = Path(sys.argv[1])
    sources = [
        Path(r"D:\trpg\大梁武侠\规则书\大梁江湖TRPG_内测第一版文件包_回填修订_2026年6月20日\01_规则书\01_玩家规则书_内测第一版.docx"),
        Path(r"D:\trpg\大梁武侠\规则书\大梁江湖TRPG_内测第一版文件包_回填修订_2026年6月20日\01_规则书\02_DM规则书_内测第一版.docx"),
        Path(r"D:\trpg\大梁武侠\规则书\大梁江湖TRPG_内测第一版文件包_回填修订_2026年6月20日\02_世界观\01_梁土百科全书_扩编总册.docx"),
        Path(r"D:\trpg\大梁武侠\规则书\大梁江湖TRPG_内测第一版文件包_回填修订_2026年6月20日\04_前置对齐材料\01_三文件对齐稿.docx"),
        Path(r"D:\trpg\大梁武侠\规则书\大梁江湖TRPG_内测第一版文件包_回填修订_2026年6月20日\04_前置对齐材料\02_目录术语冻结与排版插画流程.docx"),
        Path(r"D:\trpg\大梁武侠\规则书\大梁江湖TRPG_内测第一版文件包_回填修订_2026年6月20日\05_开发资料\01_开发者规则书_交锋辅助引擎_2026年6月20日.docx"),
    ]
    out_dir = workspace / "reports" / "database-integration-2026-07-16" / "extracted" / "docx"
    results = [extract_docx(source, out_dir) for source in sources]
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
